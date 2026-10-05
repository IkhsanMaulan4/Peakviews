"""Generate an Excel (.xlsx) peak-viewer report from PeakStore data.

Pure module: menerima dict/list biasa, tidak tahu-menahu soal PeakStore maupun
Tkinter, jadi bisa diuji standalone. Pakai openpyxl (tidak butuh Excel terinstall).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Styling constants ---
HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
HIGHLIGHT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
_THIN = Side(style="thin", color="999999")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
NUM_FMT = "#,##0"  # Excel merender pemisah ribuan sesuai locale sistem.


def _name_value_labels(show_name: bool = True) -> DataLabelList:
    """Label di bar: nama series (opsional) + nilai. Flag lain di-set False
    eksplisit karena Excel menampilkan semua yang tidak di-set."""
    labels = DataLabelList()
    labels.showSerName = show_name
    labels.showVal = True
    labels.showCatName = False
    labels.showLegendKey = False
    labels.showPercent = False
    labels.separator = "\n"
    return labels


class NoDataError(Exception):
    """Diraise saat tidak ada data peak > 0 yang bisa dilaporkan."""


def _id(n: int) -> str:
    """Format int dengan pemisah ribuan gaya Indonesia (mis. 1234 -> '1.234')."""
    return f"{n:,}".replace(",", ".")


def _nonempty_segments(peaks: dict, sources: list, segments: list) -> list:
    """Segmen (urut sesuai `segments`) yang minimal satu channel-nya peak > 0."""
    result = []
    for seg in segments:
        row = peaks.get(seg, {})
        if any((row.get(src) or 0) > 0 for src in sources):
            result.append(seg)
    return result


def _segment_total(peaks: dict, sources: list, seg: str) -> int:
    """Jumlah peak semua channel untuk satu segmen."""
    row = peaks.get(seg, {})
    return sum((row.get(src) or 0) for src in sources)


def generate_report(peaks: dict, sources: list, segments: list, out_path) -> Path:
    """Tulis laporan Excel ke out_path. Raise NoDataError bila semua segmen 0."""
    out_path = Path(out_path)
    active_segments = _nonempty_segments(peaks, sources, segments)
    if not active_segments:
        raise NoDataError("Belum ada data peak buat di-report.")

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Ringkasan"
    ws_detail = wb.create_sheet("Detail")

    _build_summary(ws_sum, peaks, sources, active_segments)
    _build_detail(ws_detail, peaks, sources, active_segments)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _build_summary(ws, peaks, sources, active_segments):
    # --- Judul + metadata ---
    ws["A1"] = "Laporan Peak Viewer"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Digenerate: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A3"] = f"Jumlah segmen ber-data: {len(active_segments)}"

    # --- Hitung angka kunci ---
    best_val, best_seg, best_src = 0, "-", "-"
    for seg in active_segments:
        row = peaks.get(seg, {})
        for src in sources:
            v = row.get(src) or 0
            if v > best_val:
                best_val, best_seg, best_src = v, seg, src

    seg_totals = {seg: _segment_total(peaks, sources, seg) for seg in active_segments}
    busiest_seg = max(seg_totals, key=seg_totals.get)

    channel_totals = {
        src: sum((peaks.get(seg, {}).get(src) or 0) for seg in active_segments)
        for src in sources
    }
    busiest_channel = max(channel_totals, key=channel_totals.get)
    n_seg = len(active_segments)
    avg_total = round(sum(seg_totals.values()) / n_seg)

    # --- Kotak highlight ---
    highlights = [
        ("Peak tertinggi keseluruhan", f"{_id(best_val)} (segmen {best_seg}, {best_src})"),
        ("Segmen paling rame", f"{busiest_seg} ({_id(seg_totals[busiest_seg])})"),
        ("Channel paling rame", f"{busiest_channel} ({_id(channel_totals[busiest_channel])})"),
        ("Rata-rata TOTAL peak antar segmen", _id(avg_total)),
    ]
    hl_title_row = 5
    ws.cell(row=hl_title_row, column=1, value="HIGHLIGHT").font = HEADER_FONT
    for i, (label, val) in enumerate(highlights, start=hl_title_row + 1):
        lc = ws.cell(row=i, column=1, value=label)
        lc.font = HEADER_FONT
        lc.fill = HIGHLIGHT_FILL
        lc.border = BORDER
        vc = ws.cell(row=i, column=2, value=val)
        vc.fill = HIGHLIGHT_FILL
        vc.border = BORDER

    # --- Tabel perbandingan channel ---
    tbl_title_row = hl_title_row + len(highlights) + 3
    ws.cell(row=tbl_title_row, column=1, value="Perbandingan Channel").font = HEADER_FONT
    hdr = tbl_title_row + 1
    for col, name in enumerate(["Channel", "Total Peak", "Rata-rata"], start=1):
        c = ws.cell(row=hdr, column=col, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
    for i, src in enumerate(sources, start=1):
        r = hdr + i
        ws.cell(row=r, column=1, value=src).border = BORDER
        tc = ws.cell(row=r, column=2, value=channel_totals[src])
        tc.number_format = NUM_FMT
        tc.border = BORDER
        ac = ws.cell(row=r, column=3, value=round(channel_totals[src] / n_seg))
        ac.number_format = NUM_FMT
        ac.border = BORDER
    last = hdr + len(sources)

    # --- Lebar kolom ---
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 14

    # --- Bar chart: total peak per channel (satu series per channel -> warna + nama sendiri) ---
    chart = BarChart()
    chart.type = "col"
    chart.title = "Total Peak per Channel"
    chart.y_axis.title = "Total Peak"
    chart.y_axis.delete = False
    chart.x_axis.delete = True  # nama channel sudah ada di label bar + legend
    for i, src in enumerate(sources, start=1):
        series = Series(Reference(ws, min_col=2, min_row=hdr + i), title=src)
        series.dLbls = _name_value_labels()
        chart.series.append(series)
    chart.height = 9
    chart.width = 16
    ws.add_chart(chart, "E5")


def _build_detail(ws, peaks, sources, active_segments):
    # --- Header ---
    headers = ["Segmen"] + list(sources) + ["TOTAL"]
    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")

    total_col = len(sources) + 2  # kolom TOTAL

    # --- Baris data ---
    for r, seg in enumerate(active_segments, start=2):
        row = peaks.get(seg, {})
        ws.cell(row=r, column=1, value=seg).border = BORDER
        for c_idx, src in enumerate(sources, start=2):
            cell = ws.cell(row=r, column=c_idx, value=(row.get(src) or 0))
            cell.number_format = NUM_FMT
            cell.border = BORDER
        tc = ws.cell(row=r, column=total_col,
                     value=_segment_total(peaks, sources, seg))
        tc.number_format = NUM_FMT
        tc.font = HEADER_FONT
        tc.border = BORDER

    last_row = 1 + len(active_segments)

    # --- Lebar kolom + freeze header ---
    ws.column_dimensions["A"].width = 34
    for c_idx in range(2, total_col + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 12
    ws.freeze_panes = "A2"

    # --- Bar chart horizontal: TOTAL peak per segmen, nama segmen di samping tiap bar ---
    chart = BarChart()
    chart.type = "bar"
    chart.varyColors = True  # tiap bar warna beda
    chart.title = "Total Peak per Segmen"
    chart.y_axis.title = "Peak Viewer"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.scaling.orientation = "maxMin"  # segmen pertama di atas
    chart.y_axis.crosses = "max"  # sumbu nilai tetap di bawah
    data = Reference(ws, min_col=total_col, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.dataLabels = _name_value_labels(show_name=False)
    chart.height = max(10, len(active_segments) * 0.9)
    chart.width = 24
    anchor = get_column_letter(total_col + 2)
    ws.add_chart(chart, f"{anchor}2")
